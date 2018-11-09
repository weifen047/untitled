#encoding:utf-8

from docx import Document


document = Document('demo.docx')

table=document.tables[0]

print table.style


print table.cell(0,0).text